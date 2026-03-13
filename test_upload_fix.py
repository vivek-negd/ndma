#!/usr/bin/env python
"""Test the fixed bulk upload with proper row filtering"""

import os
import django
import openpyxl

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'config.settings.dev')
django.setup()

from api.v1.volunteer.views import VolunteerBulkUploadAPIView

# Create test Excel with:
# Row 2: Valid volunteer (Ram Kumar Singh, Goa, South Goa)
# Row 3: Empty row (should be IGNORED)
# Row 4: Empty row (should be IGNORED)
# Row 5: Valid volunteer (Priya Sharma, Goa, South Goa)

test_file = "test_upload.xlsx"

wb = openpyxl.Workbook()
ws = wb.active

# Headers
headers = ["mis_id", "name", "mobile", "state_name", "district_name"] + [""] * 22
for col, header in enumerate(headers, 1):
    ws.cell(row=1, column=col, value=header)

# Data rows
ws['A2'] = 2001
ws['B2'] = "Ram Kumar Singh"
ws['H2'] = "9876543210"
ws['R2'] = "Goa"
ws['S2'] = "South Goa"

# Empty row 3 - should be IGNORED
ws['A3'] = ""
ws['B3'] = ""

# Empty row 4 - should be IGNORED  
ws['A4'] = ""
ws['B4'] = ""

# Valid row 5
ws['A5'] = 2002
ws['B5'] = "Priya Sharma"
ws['H5'] = "9876543211"
ws['R5'] = "Goa"
ws['S5'] = "South Goa"

wb.save(test_file)
print(f"✅ Test file created: {test_file}")

# Test parsing
print("\nTesting row parsing...")
view = VolunteerBulkUploadAPIView()

with open(test_file, 'rb') as f:
    result = view._parse_excel(f)
    print(f"Rows after parsing: {len(result)}")
    for row in result:
        print(f"  - MIS ID: {row.get('mis_id')}, Name: {row.get('name')}, State: {row.get('state_name')}, District: {row.get('district_name')}")

if len(result) == 2:
    print("\n✅ SUCCESS: Empty rows correctly filtered! Only 2 valid rows parsed.")
else:
    print(f"\n❌ FAILED: Expected 2 rows, got {len(result)}")
